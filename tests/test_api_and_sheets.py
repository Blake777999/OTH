import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.models import (
    get_all_members,
    get_member,
    Member,
    ScheduleSlotAssignment,
)
from backend.sheets import generate_csv, generate_tsv_for_sheets

client = TestClient(app)

def test_api_members_seed():
    """Verify that the 12 initial members and Jared are seeded properly."""
    res = client.get("/api/members")
    assert res.status_code == 200
    members = res.json()
    assert len(members) >= 13

    names = [m["name"] for m in members]
    assert "Dylan" in names
    assert "Nate" in names
    assert "Simon" in names
    assert "Blake Cohen" in names
    assert "Jacob" in names
    assert "Oliver" in names
    assert "Kobi" in names
    assert "Ethan" in names
    assert "Sebastian" in names
    assert "Corey" in names
    assert "Alex" in names
    assert "Blake Rosen" in names
    assert "Jared" in names

    # Jared should initially be inactive
    jared = next(m for m in members if m["name"] == "Jared")
    assert jared["active"] is False

    # Check active count is exactly 12
    active_count = sum(1 for m in members if m["active"])
    assert active_count == 12

def test_api_generate_and_retrieve_schedule():
    """Test generating a weekly schedule and fetching it back."""
    week_id = "2026-W45"
    payload = {
        "week_id": week_id,
        "use_previous_week_consistency": False,
        "min_shift_hours": 1.0,
    }
    res = client.post("/api/schedules/generate", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert len(data["assignments"]) == 280
    assert data["stats"]["fairness_score"] >= 90.0

    # Retrieve schedule
    get_res = client.get(f"/api/schedules/{week_id}")
    assert get_res.status_code == 200
    sched = get_res.json()
    assert sched["found"] is True
    assert len(sched["assignments"]) == 280

def test_csv_and_tsv_exports():
    """Verify CSV and TSV export endpoints return properly formatted grids."""
    week_id = "2026-W45"
    res_csv = client.get(f"/api/schedules/{week_id}/export.csv")
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    assert "Burn" in res_csv.text
    assert "THC" in res_csv.text

    res_tsv = client.get(f"/api/schedules/{week_id}/export.tsv")
    assert res_tsv.status_code == 200
    assert "\t" in res_tsv.text
    assert "Time Range" in res_tsv.text

def test_sheets_config_endpoints():
    """Verify getting and saving Google Sheets configuration."""
    cfg = {
        "spreadsheet_id": "test_sheet_123456",
        "sheet_name": "Weekly Schedule",
        "service_account_json": '{"type": "service_account"}',
        "auto_sync": True,
    }
    post_res = client.post("/api/sheets/config", json=cfg)
    assert post_res.status_code == 200

    get_res = client.get("/api/sheets/config")
    assert get_res.status_code == 200
    saved = get_res.json()
    assert saved["spreadsheet_id"] == "test_sheet_123456"
    assert saved["auto_sync"] is True
    assert saved["has_credentials"] is True

def test_xlsx_export_for_google_sheets():
    """Verify Excel export endpoint generates a valid multi-sheet workbook for Google Sheets."""
    import openpyxl
    import io

    week_id = "2026-W45"
    res = client.get(f"/api/schedules/{week_id}/export.xlsx")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert f'filename="old_town_hours_{week_id}.xlsx"' in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Weekly Schedule" in wb.sheetnames
    assert "Hours Summary" in wb.sheetnames

    # Check Schedule sheet
    ws_sched = wb["Weekly Schedule"]
    assert "Old Town Hours" in ws_sched["A1"].value
    assert ws_sched["A2"].value == "Time"
    assert ws_sched["B2"].value == "Monday"
    assert ws_sched["B3"].value == "Burn"
    assert ws_sched["C3"].value == "THC"

    # Check Summary sheet
    ws_sum = wb["Hours Summary"]
    assert "Member Hours Summary" in ws_sum["A1"].value
    assert ws_sum["A2"].value == "Member"
    assert ws_sum["B2"].value == "Assigned Hours"
    assert ws_sum["C2"].value == "Target Hours"
    assert ws_sum.max_row >= 14 # 12 members + header + title

def test_system_info_and_backup_restore():
    """Verify system info, export backup, and restore endpoints."""
    res_info = client.get("/api/system/info")
    assert res_info.status_code == 200
    info = res_info.json()
    assert "engine" in info
    assert "persistent" in info

    res_backup = client.get("/api/system/backup")
    assert res_backup.status_code == 200
    backup_data = res_backup.json()
    assert "members" in backup_data
    assert "master_schedules" in backup_data
    assert len(backup_data["members"]) >= 12

    res_restore = client.post("/api/system/restore", json=backup_data)
    assert res_restore.status_code == 200
    restore_res = res_restore.json()
    assert restore_res["success"] is True
    assert restore_res["counts"]["members"] >= 12

