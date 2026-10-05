import sqlite3
import json
import uuid
import os
import re
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional, Any
from pydantic import BaseModel, Field

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
DB_PATH = DATA_DIR / "scheduler.db"

DAYS_OF_WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
HOUSES = ["Burn", "THC"]
START_HOUR = 9   # 9 AM
END_HOUR = 19    # 7 PM
SLOTS_PER_DAY = 20  # 20 half-hour slots per day: 0 = 9:00-9:30, ..., 19 = 18:30-19:00

# Distinct accessible colors for members on the schedule grid
MEMBER_COLORS = [
    "#2563EB", # Blue
    "#16A34A", # Green
    "#D97706", # Amber
    "#DC2626", # Red
    "#9333EA", # Purple
    "#0D9488", # Teal
    "#EA580C", # Orange
    "#4F46E5", # Indigo
    "#059669", # Emerald
    "#DB2777", # Pink
    "#0284C7", # Sky
    "#7C3AED", # Violet
    "#65A30D", # Lime
    "#CA8A04", # Yellow
]

def slot_to_time_str(slot_idx: int) -> str:
    """Converts a slot index (0..19) to a human-readable time string like '9:00 AM'."""
    total_minutes = START_HOUR * 60 + slot_idx * 30
    hour = total_minutes // 60
    minute = total_minutes % 60
    period = "AM" if hour < 12 else "PM"
    display_hour = hour if hour <= 12 else hour - 12
    if display_hour == 0:
        display_hour = 12
    return f"{display_hour}:{minute:02d} {period}"

def slot_range_to_str(start_slot: int, end_slot: int) -> str:
    """Converts start and end slot (exclusive) to time range like '9:00 AM - 11:30 AM'."""
    start_str = slot_to_time_str(start_slot)
    end_total_minutes = START_HOUR * 60 + end_slot * 30
    end_h = end_total_minutes // 60
    end_m = end_total_minutes % 60
    period = "AM" if end_h < 12 else "PM"
    display_h = end_h if end_h <= 12 else end_h - 12
    if display_h == 0:
        display_h = 12
    end_str = f"{display_h}:{end_m:02d} {period}"
    return f"{start_str} - {end_str}"

class Member(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    email: Optional[str] = ""
    active: bool = True
    weight: float = 1.0
    shift_preference: str = "daily_short" # "daily_short" (~2h/day) or "fewer_long" (fewer days, longer blocks)
    color: str = "#2563EB"

class MasterScheduleItem(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    member_id: str
    day_of_week: int  # 0 = Monday, 6 = Sunday
    start_slot: int   # 0..19
    end_slot: int     # 1..20 (exclusive)
    label: str = "Busy"

class WeeklyOverrideItem(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    week_id: str      # e.g. "2026-W40"
    member_id: str
    day_of_week: int
    start_slot: int
    end_slot: int
    override_type: str = "busy"  # "busy" (additional conflict) or "available" (exceptionally free)
    note: str = ""

class ScheduleSlotAssignment(BaseModel):
    day_of_week: int
    slot: int
    house: str        # "Burn" or "THC"
    member_id: str
    member_name: Optional[str] = None
    color: Optional[str] = None

class GoogleSheetConfig(BaseModel):
    spreadsheet_id: str = ""
    sheet_name: str = "Weekly Schedule"
    service_account_json: Optional[str] = ""
    auto_sync: bool = False
    last_synced: Optional[str] = None

class AppCursor:
    def __init__(self, raw_cursor, is_pg: bool):
        self.raw = raw_cursor
        self.is_pg = is_pg

    def execute(self, sql: str, params=None):
        if self.is_pg:
            sql = sql.replace("?", "%s")
            if "INSERT OR REPLACE INTO app_config" in sql:
                sql = "INSERT INTO app_config (key, value) VALUES (%s, %s) ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value"
            if params is not None:
                return self.raw.execute(sql, params)
            return self.raw.execute(sql)
        else:
            if params is not None:
                return self.raw.execute(sql, params)
            return self.raw.execute(sql)

    def executemany(self, sql: str, seq_of_params):
        if self.is_pg:
            sql = sql.replace("?", "%s")
            return self.raw.executemany(sql, seq_of_params)
        return self.raw.executemany(sql, seq_of_params)

    def fetchall(self):
        return self.raw.fetchall()

    def fetchone(self):
        return self.raw.fetchone()

    def close(self):
        return self.raw.close()

class AppDBConnection:
    def __init__(self, raw_conn, is_pg: bool):
        self.raw = raw_conn
        self.is_pg = is_pg

    def cursor(self):
        if self.is_pg:
            import psycopg2.extras
            return AppCursor(self.raw.cursor(cursor_factory=psycopg2.extras.DictCursor), is_pg=True)
        return AppCursor(self.raw.cursor(), is_pg=False)

    def commit(self):
        self.raw.commit()

    def rollback(self):
        self.raw.rollback()

    def close(self):
        self.raw.close()

def get_db_connection() -> AppDBConnection:
    db_url = os.environ.get("DATABASE_URL")
    if db_url:
        if db_url.startswith("postgres://"):
            db_url = db_url.replace("postgres://", "postgresql://", 1)
        if "sslmode" not in db_url and ("supabase" in db_url or "neon.tech" in db_url or "render.com" in db_url):
            sep = "&" if "?" in db_url else "?"
            db_url = f"{db_url}{sep}sslmode=require"
        import psycopg2
        conn = psycopg2.connect(db_url)
        return AppDBConnection(conn, is_pg=True)
    else:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(str(DB_PATH))
        conn.row_factory = sqlite3.Row
        return AppDBConnection(conn, is_pg=False)

def seed_master_from_excel(cursor, excel_path: Path):
    """Seed master schedule items from avail.xlsx when initializing a fresh database."""
    try:
        import openpyxl
        column_mapping = {
            "levitt": "Nate", "nate": "Nate", "ethan": "Ethan", "kobi": "Kobi",
            "jared": "Jared", "blake rosen": "Blake Rosen", "corey": "Corey",
            "blake cohen": "Blake Cohen", "oliver": "Oliver", "seb": "Sebastian",
            "sebastian": "Sebastian", "bushy": "Alex", "alex": "Alex",
            "simon": "Simon", "dylan": "Dylan", "jacob": "Jacob"
        }
        cursor.execute("SELECT id, name FROM members")
        db_members = cursor.fetchall()
        name_to_id = {r["name"]: r["id"] for r in db_members}

        wb = openpyxl.load_workbook(excel_path)
        for day_idx, day_name in enumerate(DAYS_OF_WEEK):
            sheet_candidates = [s for s in wb.sheetnames if day_name.lower() in s.lower() or s.lower() == day_name[:3].lower()]
            if not sheet_candidates:
                continue
            ws = wb[sheet_candidates[0]]
            is_weekend = day_idx in (5, 6)
            header_row = 1
            for r in range(1, 4):
                val = str(ws.cell(r, 1).value or "").lower()
                if "free" in val or "not" in val or ws.cell(r, 2).value == "levitt":
                    header_row = r
                    break
            col_to_member = {}
            for col in range(2, ws.max_column + 1):
                raw_name = str(ws.cell(header_row, col).value or "").strip().lower()
                if raw_name in column_mapping:
                    std_name = column_mapping[raw_name]
                    if std_name in name_to_id:
                        col_to_member[col] = (name_to_id[std_name], std_name)

            for col, (mid, mname) in col_to_member.items():
                busy_slots = []
                for slot_idx in range(SLOTS_PER_DAY):
                    row_num = header_row + 1 + slot_idx
                    cell_val = ws.cell(row_num, col).value
                    has_x = str(cell_val or "").strip().upper() == "X"
                    is_busy = has_x if is_weekend else not has_x
                    if is_busy:
                        busy_slots.append(slot_idx)
                if busy_slots:
                    start = busy_slots[0]
                    prev = start
                    for s in busy_slots[1:]:
                        if s == prev + 1:
                            prev = s
                        else:
                            cursor.execute(
                                "INSERT INTO master_schedules (id, member_id, day_of_week, start_slot, end_slot, label) VALUES (?, ?, ?, ?, ?, ?)",
                                (str(uuid.uuid4()), mid, day_idx, start, prev + 1, "Class/Busy")
                            )
                            start = s
                            prev = s
                    cursor.execute(
                        "INSERT INTO master_schedules (id, member_id, day_of_week, start_slot, end_slot, label) VALUES (?, ?, ?, ?, ?, ?)",
                        (str(uuid.uuid4()), mid, day_idx, start, prev + 1, "Class/Busy")
                    )
    except Exception as e:
        print(f"Warning: could not seed master schedule from excel: {e}")

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS members (
        id TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        email TEXT,
        active INTEGER NOT NULL DEFAULT 1,
        weight REAL NOT NULL DEFAULT 1.0,
        shift_preference TEXT NOT NULL DEFAULT 'daily_short',
        color TEXT NOT NULL DEFAULT '#2563EB'
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS master_schedules (
        id TEXT PRIMARY KEY,
        member_id TEXT NOT NULL,
        day_of_week INTEGER NOT NULL,
        start_slot INTEGER NOT NULL,
        end_slot INTEGER NOT NULL,
        label TEXT NOT NULL DEFAULT 'Busy',
        FOREIGN KEY (member_id) REFERENCES members (id) ON DELETE CASCADE
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS weekly_overrides (
        id TEXT PRIMARY KEY,
        week_id TEXT NOT NULL,
        member_id TEXT NOT NULL,
        day_of_week INTEGER NOT NULL,
        start_slot INTEGER NOT NULL,
        end_slot INTEGER NOT NULL,
        override_type TEXT NOT NULL DEFAULT 'busy',
        note TEXT NOT NULL DEFAULT '',
        FOREIGN KEY (member_id) REFERENCES members (id) ON DELETE CASCADE
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS schedules (
        id TEXT PRIMARY KEY,
        week_id TEXT NOT NULL,
        day_of_week INTEGER NOT NULL,
        slot INTEGER NOT NULL,
        house TEXT NOT NULL,
        member_id TEXT NOT NULL,
        FOREIGN KEY (member_id) REFERENCES members (id) ON DELETE CASCADE,
        UNIQUE(week_id, day_of_week, slot, house)
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS app_config (
        key TEXT PRIMARY KEY,
        value TEXT NOT NULL
    )
    """)
    
    conn.commit()
    
    # Seed initial members if table is empty
    cursor.execute("SELECT COUNT(*) FROM members")
    count = cursor.fetchone()[0]
    if count == 0:
        seed_initial_members(cursor)
        conn.commit()

    # Seed master availability if empty and excel file exists
    cursor.execute("SELECT COUNT(*) FROM master_schedules")
    ms_count = cursor.fetchone()[0]
    if ms_count == 0:
        excel_path = Path(__file__).resolve().parent.parent / "avail.xlsx"
        if excel_path.exists():
            seed_master_from_excel(cursor, excel_path)
            conn.commit()
        
    conn.close()

def seed_initial_members(cursor: sqlite3.Cursor):
    initial_members = [
        {"name": "Dylan", "active": True, "pref": "daily_short"},
        {"name": "Nate", "active": True, "pref": "daily_short"},
        {"name": "Simon", "active": True, "pref": "fewer_long"},
        {"name": "Blake Cohen", "active": True, "pref": "fewer_long"},
        {"name": "Jacob", "active": True, "pref": "daily_short"},
        {"name": "Oliver", "active": True, "pref": "daily_short"},
        {"name": "Kobi", "active": True, "pref": "fewer_long"},
        {"name": "Ethan", "active": True, "pref": "daily_short"},
        {"name": "Sebastian", "active": True, "pref": "daily_short"},
        {"name": "Corey", "active": True, "pref": "fewer_long"},
        {"name": "Alex", "active": True, "pref": "daily_short"},
        {"name": "Blake Rosen", "active": True, "pref": "daily_short"},
        {"name": "Jared", "active": False, "pref": "fewer_long"}, # 13th member marked inactive until added
    ]
    for idx, m in enumerate(initial_members):
        color = MEMBER_COLORS[idx % len(MEMBER_COLORS)]
        cursor.execute(
            "INSERT INTO members (id, name, email, active, weight, shift_preference, color) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), m["name"], "", 1 if m["active"] else 0, 1.0, m["pref"], color)
        )

# DB Operations for Members
def get_all_members(only_active: bool = False) -> List[Member]:
    conn = get_db_connection()
    cursor = conn.cursor()
    if only_active:
        cursor.execute("SELECT * FROM members WHERE active = 1 ORDER BY name ASC")
    else:
        cursor.execute("SELECT * FROM members ORDER BY active DESC, name ASC")
    rows = cursor.fetchall()
    conn.close()
    return [
        Member(
            id=r["id"],
            name=r["name"],
            email=r["email"],
            active=bool(r["active"]),
            weight=r["weight"],
            shift_preference=r["shift_preference"],
            color=r["color"],
        )
        for r in rows
    ]

def get_member(member_id: str) -> Optional[Member]:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM members WHERE id = ?", (member_id,))
    row = cursor.fetchone()
    conn.close()
    if not row:
        return None
    return Member(
        id=row["id"],
        name=row["name"],
        email=row["email"],
        active=bool(row["active"]),
        weight=row["weight"],
        shift_preference=row["shift_preference"],
        color=row["color"],
    )

def create_member(member: Member) -> Member:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO members (id, name, email, active, weight, shift_preference, color) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (member.id, member.name, member.email, 1 if member.active else 0, member.weight, member.shift_preference, member.color)
    )
    conn.commit()
    conn.close()
    return member

def update_member(member: Member) -> Member:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE members SET name = ?, email = ?, active = ?, weight = ?, shift_preference = ?, color = ? WHERE id = ?",
        (member.name, member.email, 1 if member.active else 0, member.weight, member.shift_preference, member.color, member.id)
    )
    conn.commit()
    conn.close()
    return member

def delete_member(member_id: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM members WHERE id = ?", (member_id,))
    cursor.execute("DELETE FROM master_schedules WHERE member_id = ?", (member_id,))
    cursor.execute("DELETE FROM weekly_overrides WHERE member_id = ?", (member_id,))
    cursor.execute("DELETE FROM schedules WHERE member_id = ?", (member_id,))
    conn.commit()
    conn.close()

# DB Operations for Master Schedules
def get_master_schedule(member_id: Optional[str] = None) -> List[MasterScheduleItem]:
    conn = get_db_connection()
    cursor = conn.cursor()
    if member_id:
        cursor.execute("SELECT * FROM master_schedules WHERE member_id = ? ORDER BY day_of_week, start_slot", (member_id,))
    else:
        cursor.execute("SELECT * FROM master_schedules ORDER BY member_id, day_of_week, start_slot")
    rows = cursor.fetchall()
    conn.close()
    return [
        MasterScheduleItem(
            id=r["id"],
            member_id=r["member_id"],
            day_of_week=r["day_of_week"],
            start_slot=r["start_slot"],
            end_slot=r["end_slot"],
            label=r["label"],
        )
        for r in rows
    ]

def save_master_schedule_for_member(member_id: str, items: List[MasterScheduleItem]):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM master_schedules WHERE member_id = ?", (member_id,))
    for item in items:
        cursor.execute(
            "INSERT INTO master_schedules (id, member_id, day_of_week, start_slot, end_slot, label) VALUES (?, ?, ?, ?, ?, ?)",
            (item.id, member_id, item.day_of_week, item.start_slot, item.end_slot, item.label)
        )
    conn.commit()
    conn.close()

# DB Operations for Weekly Overrides
def get_weekly_overrides(week_id: str, member_id: Optional[str] = None) -> List[WeeklyOverrideItem]:
    conn = get_db_connection()
    cursor = conn.cursor()
    if member_id:
        cursor.execute("SELECT * FROM weekly_overrides WHERE week_id = ? AND member_id = ? ORDER BY day_of_week, start_slot", (week_id, member_id))
    else:
        cursor.execute("SELECT * FROM weekly_overrides WHERE week_id = ? ORDER BY member_id, day_of_week, start_slot", (week_id,))
    rows = cursor.fetchall()
    conn.close()
    return [
        WeeklyOverrideItem(
            id=r["id"],
            week_id=r["week_id"],
            member_id=r["member_id"],
            day_of_week=r["day_of_week"],
            start_slot=r["start_slot"],
            end_slot=r["end_slot"],
            override_type=r["override_type"],
            note=r["note"],
        )
        for r in rows
    ]

def save_weekly_overrides_for_member(week_id: str, member_id: str, items: List[WeeklyOverrideItem]):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM weekly_overrides WHERE week_id = ? AND member_id = ?", (week_id, member_id))
    for item in items:
        cursor.execute(
            "INSERT INTO weekly_overrides (id, week_id, member_id, day_of_week, start_slot, end_slot, override_type, note) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (item.id, week_id, member_id, item.day_of_week, item.start_slot, item.end_slot, item.override_type, item.note)
        )
    conn.commit()
    conn.close()

# DB Operations for Generated Schedules
def save_schedule(week_id: str, assignments: List[ScheduleSlotAssignment]):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM schedules WHERE week_id = ?", (week_id,))
    for a in assignments:
        cursor.execute(
            "INSERT INTO schedules (id, week_id, day_of_week, slot, house, member_id) VALUES (?, ?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), week_id, a.day_of_week, a.slot, a.house, a.member_id)
        )
    conn.commit()
    conn.close()

def get_schedule(week_id: str) -> List[ScheduleSlotAssignment]:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT s.day_of_week, s.slot, s.house, s.member_id,
               COALESCE(m.name, '❌ X (UNFILLED)') as member_name,
               COALESCE(m.color, '#DC2626') as color
        FROM schedules s
        LEFT JOIN members m ON s.member_id = m.id
        WHERE s.week_id = ?
        ORDER BY s.day_of_week, s.slot, s.house
    """, (week_id,))
    rows = cursor.fetchall()
    conn.close()
    return [
        ScheduleSlotAssignment(
            day_of_week=r["day_of_week"],
            slot=r["slot"],
            house=r["house"],
            member_id=r["member_id"],
            member_name=r["member_name"],
            color=r["color"],
        )
        for r in rows
    ]

def get_all_scheduled_weeks() -> List[str]:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT week_id FROM schedules ORDER BY week_id DESC")
    rows = cursor.fetchall()
    conn.close()
    return [r[0] for r in rows]

# Config Operations
def get_config(key: str, default: str = "") -> str:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM app_config WHERE key = ?", (key,))
    row = cursor.fetchone()
    conn.close()
    return row[0] if row else default

def set_config(key: str, value: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO app_config (key, value) VALUES (?, ?)", (key, value))
    conn.commit()
    conn.close()

# System & Backup Operations
def get_db_info() -> Dict[str, Any]:
    db_url = os.environ.get("DATABASE_URL")
    if db_url:
        import re
        safe_url = re.sub(r':([^@]+)@', ':****@', db_url)
        return {
            "engine": "postgresql",
            "persistent": True,
            "display": "Cloud PostgreSQL (Persistent across spin-ups & upgrades)",
            "safe_url": safe_url,
        }
    return {
        "engine": "sqlite",
        "persistent": False,
        "display": "Local SQLite (Ephemeral on Render free tier)",
        "path": str(DB_PATH),
    }

def export_full_backup() -> Dict[str, Any]:
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT id, name, email, active, weight, shift_preference, color FROM members ORDER BY name")
    members = [
        {
            "id": r["id"],
            "name": r["name"],
            "email": r["email"] or "",
            "active": bool(r["active"]),
            "weight": float(r["weight"]),
            "shift_preference": r["shift_preference"],
            "color": r["color"],
        }
        for r in cursor.fetchall()
    ]

    cursor.execute("SELECT id, member_id, day_of_week, start_slot, end_slot, label FROM master_schedules ORDER BY day_of_week, start_slot")
    master_schedules = [
        {
            "id": r["id"],
            "member_id": r["member_id"],
            "day_of_week": int(r["day_of_week"]),
            "start_slot": int(r["start_slot"]),
            "end_slot": int(r["end_slot"]),
            "label": r["label"] or "Busy",
        }
        for r in cursor.fetchall()
    ]

    cursor.execute("SELECT id, week_id, member_id, day_of_week, start_slot, end_slot, override_type, note FROM weekly_overrides ORDER BY week_id, day_of_week, start_slot")
    weekly_overrides = [
        {
            "id": r["id"],
            "week_id": r["week_id"],
            "member_id": r["member_id"],
            "day_of_week": int(r["day_of_week"]),
            "start_slot": int(r["start_slot"]),
            "end_slot": int(r["end_slot"]),
            "override_type": r["override_type"],
            "note": r["note"] or "",
        }
        for r in cursor.fetchall()
    ]

    cursor.execute("SELECT id, week_id, day_of_week, slot, house, member_id FROM schedules ORDER BY week_id, day_of_week, slot, house")
    schedules = [
        {
            "id": r["id"],
            "week_id": r["week_id"],
            "day_of_week": int(r["day_of_week"]),
            "slot": int(r["slot"]),
            "house": r["house"],
            "member_id": r["member_id"],
        }
        for r in cursor.fetchall()
    ]

    cursor.execute("SELECT key, value FROM app_config")
    app_config = {r["key"]: r["value"] for r in cursor.fetchall()}
    conn.close()

    return {
        "app": "Old Town Hours",
        "version": 1,
        "exported_at": datetime.now().isoformat(),
        "counts": {
            "members": len(members),
            "master_schedules": len(master_schedules),
            "weekly_overrides": len(weekly_overrides),
            "schedules": len(schedules),
        },
        "members": members,
        "master_schedules": master_schedules,
        "weekly_overrides": weekly_overrides,
        "schedules": schedules,
        "app_config": app_config,
    }

def import_full_backup(data: Dict[str, Any]) -> Dict[str, Any]:
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM schedules")
        cursor.execute("DELETE FROM weekly_overrides")
        cursor.execute("DELETE FROM master_schedules")
        cursor.execute("DELETE FROM members")
        cursor.execute("DELETE FROM app_config")

        for m in data.get("members", []):
            cursor.execute(
                "INSERT INTO members (id, name, email, active, weight, shift_preference, color) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (m["id"], m["name"], m.get("email", ""), 1 if m.get("active") else 0, m.get("weight", 1.0), m.get("shift_preference", "daily_short"), m.get("color", "#2563EB"))
            )
        for ms in data.get("master_schedules", []):
            cursor.execute(
                "INSERT INTO master_schedules (id, member_id, day_of_week, start_slot, end_slot, label) VALUES (?, ?, ?, ?, ?, ?)",
                (ms["id"], ms["member_id"], ms["day_of_week"], ms["start_slot"], ms["end_slot"], ms.get("label", "Busy"))
            )
        for wo in data.get("weekly_overrides", []):
            cursor.execute(
                "INSERT INTO weekly_overrides (id, week_id, member_id, day_of_week, start_slot, end_slot, override_type, note) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (wo["id"], wo["week_id"], wo["member_id"], wo["day_of_week"], wo["start_slot"], wo["end_slot"], wo.get("override_type", "busy"), wo.get("note", ""))
            )
        for s in data.get("schedules", []):
            cursor.execute(
                "INSERT INTO schedules (id, week_id, day_of_week, slot, house, member_id) VALUES (?, ?, ?, ?, ?, ?)",
                (s.get("id", str(uuid.uuid4())), s["week_id"], s["day_of_week"], s["slot"], s["house"], s["member_id"])
            )
        for k, v in data.get("app_config", {}).items():
            cursor.execute("INSERT OR REPLACE INTO app_config (key, value) VALUES (?, ?)", (k, v))

        conn.commit()
        return {
            "success": True,
            "message": f"Successfully restored {len(data.get('members', []))} members, {len(data.get('master_schedules', []))} master blocks, {len(data.get('weekly_overrides', []))} overrides, {len(data.get('schedules', []))} slot assignments.",
            "counts": {
                "members": len(data.get("members", [])),
                "master_schedules": len(data.get("master_schedules", [])),
                "weekly_overrides": len(data.get("weekly_overrides", [])),
                "schedules": len(data.get("schedules", [])),
            }
        }
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        conn.close()

