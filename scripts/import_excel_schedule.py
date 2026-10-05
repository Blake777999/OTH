import os
import uuid
import sqlite3
import openpyxl
from pathlib import Path

# Paths
BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "data" / "scheduler.db"
EXCEL_PATH = BASE_DIR / "avail.xlsx"

DAYS_OF_WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# Column header normalization mapping
COLUMN_MAPPING = {
    "levitt": "Nate",
    "nate": "Nate",
    "ethan": "Ethan",
    "kobi": "Kobi",
    "jared": "Jared",
    "blake rosen": "Blake Rosen",
    "corey": "Corey",
    "blake cohen": "Blake Cohen",
    "oliver": "Oliver",
    "seb": "Sebastian",
    "sebastian": "Sebastian",
    "bushy": "Alex",
    "alex": "Alex",
    "simon": "Simon",
    "dylan": "Dylan",
    "jacob": "Jacob"
}

def import_availability():
    if not EXCEL_PATH.exists():
        print(f"Error: {EXCEL_PATH} not found!")
        return False

    from backend.models import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()

    # Load member database mapping
    cursor.execute("SELECT id, name FROM members")
    db_members = dict(cursor.fetchall())
    name_to_id = {name: mid for mid, name in db_members.items()}

    wb = openpyxl.load_workbook(EXCEL_PATH)
    master_items = []
    member_busy_counts = {name: 0 for name in db_members.values()}
    member_free_counts = {name: 0 for name in db_members.values()}

    for day_idx, day_name in enumerate(DAYS_OF_WEEK):
        if day_name not in wb.sheetnames:
            print(f"Warning: Sheet {day_name} not found in workbook!")
            continue

        ws = wb[day_name]
        is_weekend = day_idx in (5, 6) # Saturday (5), Sunday (6)

        # Find header row containing the member names
        header_row = 1
        for r in range(1, 4):
            val = str(ws.cell(r, 1).value or "").lower()
            if "free" in val or "not" in val or ws.cell(r, 2).value == "levitt":
                header_row = r
                break

        col_to_member = {}
        for col in range(2, ws.max_column + 1):
            raw_name = str(ws.cell(header_row, col).value or "").strip().lower()
            if raw_name in COLUMN_MAPPING:
                std_name = COLUMN_MAPPING[raw_name]
                if std_name in name_to_id:
                    col_to_member[col] = (name_to_id[std_name], std_name)

        # Process 20 slots (9am to 7pm)
        for col, (mid, mname) in col_to_member.items():
            busy_slots = []
            for slot_idx in range(20):
                row_num = header_row + 1 + slot_idx
                cell_val = ws.cell(row_num, col).value
                has_x = str(cell_val or "").strip().upper() == "X"

                if not is_weekend:
                    # Weekday: X = Free, Blank = Busy
                    is_busy = not has_x
                else:
                    # Weekend: X = Busy, Blank = Free
                    is_busy = has_x

                if is_busy:
                    busy_slots.append(slot_idx)
                    member_busy_counts[mname] += 1
                else:
                    member_free_counts[mname] += 1

            # Merge contiguous busy slots for clean DB records
            if busy_slots:
                start = busy_slots[0]
                prev = start
                for s in busy_slots[1:]:
                    if s == prev + 1:
                        prev = s
                    else:
                        master_items.append((
                            str(uuid.uuid4()),
                            mid,
                            day_idx,
                            start,
                            prev + 1,
                            "Class/Busy"
                        ))
                        start = s
                        prev = s
                master_items.append((
                    str(uuid.uuid4()),
                    mid,
                    day_idx,
                    start,
                    prev + 1,
                    "Class/Busy"
                ))

    # Replace existing master schedules
    print("Replacing master schedules in database...")
    cursor.execute("DELETE FROM master_schedules")
    cursor.executemany(
        "INSERT INTO master_schedules (id, member_id, day_of_week, start_slot, end_slot, label) VALUES (?, ?, ?, ?, ?, ?)",
        master_items
    )

    # Clear old schedules so the user can generate a fresh one
    cursor.execute("DELETE FROM schedules")

    conn.commit()
    conn.close()

    print(f"Successfully imported {len(master_items)} master schedule blocks across all 7 days!")
    print("\nMember Weekly Availability Breakdown:")
    for name in sorted(member_free_counts.keys()):
        free_s = member_free_counts[name]
        busy_s = member_busy_counts[name]
        print(f"  {name:12s}: {free_s:3d} free slots ({free_s * 0.5:4.1f}h) | {busy_s:3d} busy slots ({busy_s * 0.5:4.1f}h)")

    return True

if __name__ == "__main__":
    import_availability()
